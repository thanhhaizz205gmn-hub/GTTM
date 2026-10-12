import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def format_row(row, is_header=False):
    for cell in row.cells:
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        if is_header:
            set_cell_background(cell, "1F4E79")  # Dark Tech Blue
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(12)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
        else:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(12)
                    run.font.color.rgb = RGBColor(50, 50, 50)

doc = docx.Document()

# Set standard margins (1 inch)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Helper for Headings
def add_custom_heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)

# 1. DANH MỤC TỪ VIẾT TẮT
add_custom_heading("DANH MỤC TỪ VIẾT TẮT")

abbreviations_data = [
    ("AI", "Artificial Intelligence — Trí tuệ nhân tạo"),
    ("API", "Application Programming Interface — Giao diện lập trình ứng dụng"),
    ("BMS", "Building Management System — Hệ thống quản lý tòa nhà"),
    ("CV", "Cross-Validation — Kiểm chứng chéo"),
    ("EDA", "Exploratory Data Analysis — Phân tích khám phá dữ liệu"),
    ("FastAPI", "Fast Application Programming Interface Framework — Framework Python phát triển RESTful API hiệu năng cao"),
    ("HVAC", "Heating, Ventilation, and Air Conditioning — Hệ thống sưởi, thông gió và điều hòa không khí"),
    ("IoT", "Internet of Things — Internet kết nối mọi vật"),
    ("LightGBM", "Light Gradient Boosting Machine — Thuật toán học máy tăng cường độ dốc tối ưu bộ nhớ"),
    ("MAE", "Mean Absolute Error — Sai số tuyệt đối trung bình"),
    ("ML", "Machine Learning — Học máy"),
    ("n8n", "Workflow Automation Platform — Nền tảng tự động hóa luồng dữ liệu"),
    ("pax", "Passenger — Hành khách (đơn vị đo mật độ người)"),
    ("REST", "Representational State Transfer — Kiểu kiến trúc dịch vụ web"),
    ("RH", "Relative Humidity — Độ ẩm tương đối (%)"),
    ("RMSE", "Root Mean Square Error — Căn bậc hai sai số bình phương trung bình"),
    ("R²", "Coefficient of Determination — Hệ số xác định (đánh giá độ phù hợp mô hình)"),
    ("UI/UX", "User Interface / User Experience — Giao diện / Trải nghiệm người dùng"),
    ("Unity 3D", "Unity Real-Time Development Platform — Engine đồ họa trực quan hóa 3D"),
    ("WebGL", "Web Graphics Library — Thư viện chuẩn đồ họa 3D trên trình duyệt web")
]

table1 = doc.add_table(rows=1, cols=2)
table1.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells = table1.rows[0].cells
hdr_cells[0].text = "Từ viết tắt"
hdr_cells[1].text = "Tên tiếng Anh & Giải nghĩa tiếng Việt"
format_row(table1.rows[0], is_header=True)

for i, (abbr, desc) in enumerate(abbreviations_data):
    row_cells = table1.add_row().cells
    row_cells[0].text = abbr
    row_cells[1].text = desc
    format_row(table1.rows[-1], is_header=False)
    if i % 2 == 1:
        set_cell_background(row_cells[0], "F9FBFD")
        set_cell_background(row_cells[1], "F9FBFD")

# Set widths for Table 1
for row in table1.rows:
    row.cells[0].width = Inches(1.8)
    row.cells[1].width = Inches(4.7)

# 2. DANH MỤC HÌNH VẼ
add_custom_heading("DANH MỤC HÌNH VẼ")

figures_data = [
    ("Hình 1.1", "Sơ đồ tổng quan hệ thống Digital Twin ứng dụng trong ga tàu điện ngầm"),
    ("Hình 2.1", "Biểu đồ cân bằng tải nhiệt HVAC trong không gian ga tàu ngầm"),
    ("Hình 2.2", "Sơ đồ Kiến trúc 4 lớp tích hợp (Four-Layer Integrated Digital Twin Architecture)"),
    ("Hình 2.3", "Sơ đồ luồng Pipeline hoạt động tổng thể của hệ thống"),
    ("Hình 2.4", "Quy trình 9 bước huấn luyện và tinh chỉnh mô hình Machine Learning"),
    ("Hình 3.1", "Biểu đồ tương quan Pearson giữa thời tiết, mật độ hành khách và tải HVAC"),
    ("Hình 3.2", "Giao diện trực quan hóa 3D WebGL không gian ga tàu ngầm trên Unity"),
    ("Hình 3.3", "Biểu đồ Heatmap nhiệt độ 3D động hiển thị phân bố nhiệt thời gian thực"),
    ("Hình 4.1", "Biểu đồ so sánh hiệu năng giữa các thuật toán Machine Learning (R², RMSE)"),
    ("Hình 4.2", "Biểu đồ phân bố sai số dự báo trước và sau khi tối ưu Hyperparameter"),
    ("Hình 4.3", "Giao diện bảng điều khiển Dashboard và cảnh báo tải cao (High-load Alert)")
]

table2 = doc.add_table(rows=1, cols=2)
table2.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells2 = table2.rows[0].cells
hdr_cells2[0].text = "Ký hiệu hình"
hdr_cells2[1].text = "Tên hình vẽ"
format_row(table2.rows[0], is_header=True)

for i, (fig_id, fig_title) in enumerate(figures_data):
    row_cells = table2.add_row().cells
    row_cells[0].text = fig_id
    row_cells[1].text = fig_title
    format_row(table2.rows[-1], is_header=False)
    if i % 2 == 1:
        set_cell_background(row_cells[0], "F9FBFD")
        set_cell_background(row_cells[1], "F9FBFD")

# Set widths for Table 2
for row in table2.rows:
    row.cells[0].width = Inches(1.8)
    row.cells[1].width = Inches(4.7)

# 3. DANH MỤC BẢNG BIỂU
add_custom_heading("DANH MỤC BẢNG BIỂU")

tables_data = [
    ("Bảng 1.1", "So sánh giải pháp Digital Twin với các hệ thống BMS truyền thống"),
    ("Bảng 2.1", "Danh sách các tham số đầu vào và đầu ra của mô hình dự báo tải nhiệt"),
    ("Bảng 2.2", "Cấu trúc các bảng dữ liệu lưu trữ trên Supabase (PostgreSQL Cloud)"),
    ("Bảng 3.1", "Thống kê mô tả tập dữ liệu vận hành ga tàu (hanoi_station_raw.csv)"),
    ("Bảng 3.2", "Thông số cấu hình môi trường thử nghiệm và hạ tầng triển khai"),
    ("Bảng 4.1", "Kết quả so sánh hiệu năng 10 thuật toán Machine Learning cơ sở (Baseline)"),
    ("Bảng 4.2", "Kết quả tối ưu tham số mô hình LightGBM bằng RandomizedSearchCV"),
    ("Bảng 4.3", "Đánh giá độ ổn định mô hình bằng phương pháp Monte Carlo Bootstrap (100 lần)"),
    ("Bảng 4.4", "Kết quả đo kiểm độ trễ API và hiệu năng thời gian thực trên Unity 3D WebGL")
]

table3 = doc.add_table(rows=1, cols=2)
table3.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells3 = table3.rows[0].cells
hdr_cells3[0].text = "Ký hiệu bảng"
hdr_cells3[1].text = "Tên bảng biểu"
format_row(table3.rows[0], is_header=True)

for i, (tbl_id, tbl_title) in enumerate(tables_data):
    row_cells = table3.add_row().cells
    row_cells[0].text = tbl_id
    row_cells[1].text = tbl_title
    format_row(table3.rows[-1], is_header=False)
    if i % 2 == 1:
        set_cell_background(row_cells[0], "F9FBFD")
        set_cell_background(row_cells[1], "F9FBFD")

# Set widths for Table 3
for row in table3.rows:
    row.cells[0].width = Inches(1.8)
    row.cells[1].width = Inches(4.7)

output_path = r"d:\Đồ Án GTTM\Danh_Muc_Tu_Viet_Tat_Va_Bang_Bieu.docx"
doc.save(output_path)

