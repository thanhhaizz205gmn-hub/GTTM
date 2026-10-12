import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def format_row(row, is_header=False):
    for cell in row.cells:
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        if is_header:
            set_cell_background(cell, "1F4E79")
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(11)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
        else:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(11)
                    run.font.color.rgb = RGBColor(50, 50, 50)

doc = docx.Document()

# Set standard 1-inch margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Style functions
def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)

def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.italic = True
    run.font.color.rgb = RGBColor(80, 80, 80)

def add_h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)

def add_h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(40, 40, 40)

def add_p(text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 30, 30)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.italic = italic
    run.font.color.rgb = RGBColor(40, 40, 40)
    return p

def add_code_block(code_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    
    # Table wrapping for background
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F2F4F7")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_after = Pt(0)
    run = cp.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(30, 30, 30)

def add_bullet(text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)

# --- DOCUMENT CONTENT ---

add_title("HƯỚNG DẪN CHI TIẾT DEPLOY & PUBLIC DỰ ÁN WEB / DIGITAL TWIN LÊN INTERNET")
add_subtitle("Quy trình chuẩn từ Mua Tên Miền Hostinger, Cấu hình Cloudflare Zero Trust Tunnel đến Docker Orchestration")

# CHƯƠNG 1
add_h1("CHƯƠNG 1: TỔNG QUAN KIẾN TRÚC TRIỂN KHAI")
add_p("Tài liệu này hướng dẫn chi tiết từng bước đưa một dự án Web / Digital Twin (gồm Backend API, Nginx Web Server, 3D WebGL và Bộ mô phỏng IoT) ra Internet với Tên miền riêng có chứng chỉ mã hóa SSL/HTTPS an toàn 100%.")
add_p("Mô hình triển khai áp dụng kiến trúc Zero Trust tiên tiến nhất hiện nay:", bold_prefix="Mô hình triển khai: ")

add_bullet("Người dùng truy cập domain riêng mua tại Hostinger qua giao thức bảo mật HTTPS (Port 443).", bold_prefix="1. Người dùng công cộng: ")
add_bullet("Hệ thống DNS và mạng lưới bảo vệ toàn cầu của Cloudflare (chống DDoS, mã hóa SSL tự động).", bold_prefix="2. Cloudflare Edge Network: ")
add_bullet("Kênh kết nối bảo mật mã hóa 2 chiều kết nối trực tiếp từ Cloudflare về máy chủ/máy trạm của bạn mà KHÔNG CẦN mở Port Router (No Port Forwarding), KHÔNG CẦN IP Tĩnh (No Static IP).", bold_prefix="3. Cloudflare Tunnel (cloudflared): ")
add_bullet("4 Container đóng gói độc lập bao gồm Backend FastAPI (Port 8000), Nginx Reverse Proxy (Port 80), IoT Simulator và Cloudflared Daemon.", bold_prefix="4. Môi trường Docker Container: ")

# CHƯƠNG 2
add_h1("CHƯƠNG 2: MUA VÀ QUẢN LÝ TÊN MIỀN TRÊN HOSTINGER")
add_h2("Bước 2.1: Mua Tên miền (Domain)")
add_bullet("Truy cập trang chủ Hostinger (hostinger.vn hoặc hostinger.com) và đăng ký tài khoản.", bold_prefix="1. Đăng ký tài khoản: ")
add_bullet("Vào mục Tên miền (Domains) -> Tìm kiếm tên miền mong muốn (Ví dụ: digitaltwin-hanoi.com hoặc yourproject.site).", bold_prefix="2. Tìm kiếm tên miền: ")
add_bullet("Tiến hành thanh toán và hoàn tất xác minh Email sở hữu tên miền theo yêu cầu của Hostinger.", bold_prefix="3. Thanh toán: ")

add_h2("Bước 2.2: Đăng nhập Trang quản trị Hostinger")
add_p("Sau khi mua thành công, đăng nhập vào hPanel của Hostinger -> Chọn mục Domains (Tên miền) -> Chọn tên miền vừa mua để chuẩn bị trỏ NameServers sang Cloudflare.")

# CHƯƠNG 3
add_h1("CHƯƠNG 3: CẤU HÌNH CLOUDFLARE & TRỎ DNS NAMESERVERS")
add_h2("Bước 3.1: Đăng ký Cloudflare và Thêm Website (Add a Site)")
add_bullet("Tạo tài khoản miễn phí tại dash.cloudflare.com.", bold_prefix="1. Tạo tài khoản: ")
add_bullet("Nhấn nút Add a Site -> Nhập tên miền bạn vừa mua tại Hostinger -> Chọn gói Free (Miễn phí).", bold_prefix="2. Thêm tên miền: ")
add_bullet("Cloudflare sẽ quét bản ghi DNS cũ và cung cấp cho bạn 2 bản ghi NameServers mới (Ví dụ: alan.ns.cloudflare.com và amir.ns.cloudflare.com).", bold_prefix="3. Lấy NameServers: ")

add_h2("Bước 3.2: Đổi NameServers tại Hostinger")
add_bullet("Quay lại hPanel của Hostinger -> Vào quản lý Tên miền -> Chọn Quản lý DNS / NameServers (DNS / Nameservers).", bold_prefix="1. Vào Hostinger: ")
add_bullet("Chọn Thay đổi NameServers (Change Nameservers) -> Chọn Thay đổi thủ công (Use custom nameservers).", bold_prefix="2. Thay đổi thủ công: ")
add_bullet("Điền 2 địa chỉ NameServers mà Cloudflare cung cấp vào ô NS1 và NS2 -> Nhấn Lưu (Save).", bold_prefix="3. Cập nhật: ")
add_bullet("Chờ từ 5 - 15 phút để DNS toàn cầu cập nhật. Cloudflare sẽ gửi Email thông báo 'Website is now active'.", bold_prefix="4. Chờ kích hoạt: ")

add_h2("Bước 3.3: Tạo Cloudflare Tunnel (Zero Trust)")
add_bullet("Tại bảng điều khiển Cloudflare, chọn menu Zero Trust (ở cột bên trái).", bold_prefix="1. Vào Zero Trust: ")
add_bullet("Chọn Networks -> Tunnels -> Nhấn nút Create a Tunnel.", bold_prefix="2. Tạo Tunnel: ")
add_bullet("Chọn loại Connector là Cloudflared -> Đặt tên cho Tunnel (Ví dụ: digital_twin_tunnel) -> Nhấn Save Tunnel.", bold_prefix="3. Đặt tên: ")

add_h2("Bước 3.4: Lấy mã Cloudflare Tunnel Token")
add_p("Tại màn hình Install and run a connector, chọn tab Docker. Bạn sẽ thấy một đoạn mã token dạng:")
add_code_block("eyJhIjoiYjgyNDYzN2... (đoạn token dài chứa thông tin xác thực tunnel)")
add_p("Lưu trữ đoạn Token này cẩn thận để điền vào file docker-compose.yml ở Chương 4.")

add_h2("Bước 3.5: Cấu hình Public Hostname (Routing)")
add_bullet("Tại bước tiếp theo (Public Hostnames), nhấn Add a public hostname.", bold_prefix="1. Thêm Hostname: ")
add_bullet("Subdomain: Để trống (hoặc điền 'api' nếu làm subdomain), Domain: Chọn tên miền của bạn.", bold_prefix="2. Tên miền: ")
add_bullet("Type: HTTP, URL: nginx:80 (hoặc host.docker.internal:8000 tùy cấu hình).", bold_prefix="3. Service Target: ")
add_bullet("Nhấn Save Hostname. Cloudflare sẽ tự động tạo bản ghi CNAME trỏ tên miền về Tunnel an toàn.", bold_prefix="4. Hoàn tất: ")

# CHƯƠNG 4
add_h1("CHƯƠNG 4: CẤU HÌNH DOCKER & DOCKER COMPOSE")
add_h2("Bước 4.1: Tạo Dockerfile cho Backend App")
add_p("Tạo file Dockerfile tại thư mục gốc của dự án:")
add_code_block("""# Sử dụng hình ảnh Python 3.11 chính thức
FROM python:3.11-slim

WORKDIR /app

# Cấu hình múi giờ Việt Nam (GMT+7)
ENV TZ=Asia/Ho_Chi_Minh

# Cài đặt các thư viện hệ thống cần thiết cho C++/LightGBM
RUN apt-get update && apt-get install -y --no-install-recommends \\
    build-essential \\
    libgomp1 \\
    curl \\
    tzdata \\
    && rm -rf /var/lib/apt/lists/*

# Cài đặt Python Dependencies
COPY requirements.txt .
RUN pip install --no-cache-dir -r requirements.txt

# Copy mã nguồn và mô hình học máy
COPY app.py .
COPY digital_twin_hvac_model.pkl .
COPY mock_iot_sensor.py .
COPY templates/ ./templates/
COPY Build_Web/ ./Build_Web/

EXPOSE 8000

CMD ["uvicorn", "app:app", "--host", "0.0.0.0", "--port", "8000"]""")

add_h2("Bước 4.2: Cấu hình Nginx Reverse Proxy (nginx.conf)")
add_p("Tạo file nginx.conf để xử lý file tĩnh 3D WebGL (kèm chuẩn nén Gzip) và định tuyến API:")
add_code_block("""server {
    listen 80;
    server_name localhost;

    location / {
        proxy_pass http://app:8000;
        proxy_set_header Host $host;
        proxy_set_header X-Real-IP $remote_addr;
        proxy_set_header X-Forwarded-For $proxy_add_x_forwarded_for;
        proxy_set_header X-Forwarded-Proto $scheme;
    }

    # Hỗ trợ nén Gzip cho Unity WebGL Build Files
    location ~ \\.(data|wasm|js)\\.gz$ {
        root /usr/share/nginx/html;
        add_header Content-Encoding gzip;
        default_type application/octet-stream;
    }
}""")

add_h2("Bước 4.3: Tạo docker-compose.yml Tích Hợp 4 Services")
add_p("Tạo file docker-compose.yml liên kết tất cả các thành phần:")
add_code_block("""version: '3.8'

services:
  app:
    build:
      context: .
      dockerfile: Dockerfile
    container_name: digital_twin_app
    ports:
      - "8000:8000"
    restart: always
    environment:
      - PYTHONUNBUFFERED=1
      - TZ=Asia/Ho_Chi_Minh

  mock_sensor:
    build:
      context: .
      dockerfile: Dockerfile
    container_name: digital_twin_mock_sensor
    command: ["python", "mock_iot_sensor.py"]
    depends_on:
      - app
    restart: always
    environment:
      - PYTHONUNBUFFERED=1
      - TZ=Asia/Ho_Chi_Minh

  nginx:
    image: nginx:alpine
    container_name: digital_twin_nginx
    ports:
      - "80:80"
    volumes:
      - ./nginx.conf:/etc/nginx/conf.d/default.conf:ro
      - ./Build_Web:/usr/share/nginx/html:ro
    depends_on:
      - app
    restart: always

  cloudflared:
    image: cloudflare/cloudflared:latest
    container_name: digital_twin_cloudflared
    command: tunnel --no-autoupdate run --token CHÈN_TOKEN_CLOUDFLARE_CỦA_BẠN_VÀO_ĐÂY
    extra_hosts:
      - "host.docker.internal:host-gateway"
    depends_on:
      - app
      - nginx
    restart: always""")

# CHƯƠNG 5
add_h1("CHƯƠNG 5: KHỞI CHẠY VÀ KIỂM TRA HỆ THỐNG")
add_h2("Bước 5.1: Chạy lệnh Rebuild và Start Docker")
add_p("Mở terminal tại thư mục dự án và chạy lệnh:")
add_code_block("docker-compose up -d --build")

add_h2("Bước 5.2: Kiểm tra trạng thái hệ thống")
add_p("Kiểm tra 4 container đã hoạt động bình thường (trạng thái Up):")
add_code_block("docker ps")
add_p("Xem log kiểm tra kết nối Cloudflare Tunnel:")
add_code_block("docker-compose logs -f cloudflared")

add_h2("Bước 5.3: Truy cập Tên miền")
add_p("Mở trình duyệt bất kỳ (trên máy tính hoặc điện thoại) và truy cập tên miền của bạn:")
add_code_block("https://tendoan-cuaban.com")
add_p("Bạn sẽ thấy biểu tượng ổ khóa bảo mật HTTPS màu xanh cùng toàn bộ giao diện 3D WebGL và Dashboard chạy mượt mà!")

# CHƯƠNG 6
add_h1("CHƯƠNG 6: CÁC LỖI THƯỜNG GẶP VÀ CÁCH XỬ LÝ (TROUBLESHOOTING)")

add_bullet("Nguyên nhân: Tunnel token sai hoặc service Target trong Cloudflare đặt chưa đúng tên container (Phải là http://nginx:80 hoặc http://app:8000).", bold_prefix="1. Lỗi 502 Bad Gateway: ")
add_bullet("Khắc phục: Trong Dockerfile và docker-compose.yml phải thiết lập ENV TZ=Asia/Ho_Chi_Minh và mã nguồn Python dùng datetime.now(timezone(timedelta(hours=7))).", bold_prefix="2. Lỗi lệch múi giờ trên Biểu đồ: ")
add_bullet("Khắc phục: Kiểm tra cấu hình Nginx MIME types và thêm header Content-Encoding gzip cho các file .data.gz và .wasm.gz.", bold_prefix="3. Lỗi không tải được mô hình 3D Unity WebGL: ")

output_path = r"d:\Đồ Án GTTM\Huong_Dan_Trien_Khai_Domain_Cloudflare_Docker.docx"
doc.save(output_path)
